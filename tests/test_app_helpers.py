from datetime import date, datetime

import pytest
from flask import session

import app as app_module
from app import (
    Auditoria,
    MovimientoPuntos,
    Notificacion,
    Permiso,
    Rol,
    Usuario,
    build_filtered_auditoria_query,
    build_filtered_movimientos_query,
    build_report_context,
    calcular_saldo_local,
    crear_base,
    exportar_reportes_csv,
    exportar_reportes_excel,
    get_db,
    get_report_filters,
    get_unread_notifications_count,
    get_usuario,
    log_auditoria,
    mapear_rol_por_grupo,
    obtener_pagina_inicio,
    parse_date_param,
    requiere_permiso,
    serialize_auditoria,
    serialize_datetime,
    serialize_movimiento,
    ya_tiene_resena_local,
)


# ---------------------------------------------------------------- serializers


def test_serialize_datetime_formats_datetime():
    assert serialize_datetime(datetime(2024, 5, 4, 9, 8, 7)) == "2024-05-04 09:08:07"


def test_serialize_datetime_none_is_empty_string():
    assert serialize_datetime(None) == ""


def test_serialize_datetime_other_types_use_str():
    assert serialize_datetime(date(2024, 5, 4)) == "2024-05-04"
    assert serialize_datetime(15) == "15"


def test_serialize_movimiento_normaliza_nulos():
    movimiento = MovimientoPuntos(
        id=1,
        paciente_id=99,
        idcita=None,
        tipo=None,
        cantidad=None,
        usuario_admin=None,
        fecha=datetime(2024, 1, 2, 3, 4, 5),
        motivo=None,
    )

    assert serialize_movimiento(movimiento) == {
        "id": 1,
        "paciente_id": "99",
        "idcita": "",
        "tipo": "",
        "cantidad": 0,
        "usuario_admin": "",
        "fecha": "2024-01-02 03:04:05",
        "motivo": "",
    }


def test_serialize_auditoria_normaliza_nulos():
    registro = Auditoria(
        id=7,
        usuario_admin="admin",
        accion=None,
        usuario_afectado=None,
        detalle=None,
        fecha=None,
    )

    assert serialize_auditoria(registro) == {
        "id": 7,
        "usuario_admin": "admin",
        "accion": "",
        "usuario_afectado": "",
        "detalle": "",
        "fecha": "",
    }


# ---------------------------------------------------------------------- filtros


@pytest.mark.parametrize("valor", ["", None, "no-es-fecha", "2024-13-01"])
def test_parse_date_param_invalidos(valor):
    assert parse_date_param(valor) is None


def test_parse_date_param_inicio_de_dia():
    assert parse_date_param("2024-03-10") == datetime(2024, 3, 10, 0, 0, 0)


def test_parse_date_param_fin_de_dia():
    assert parse_date_param("2024-03-10", end_of_day=True) == datetime(
        2024, 3, 10, 23, 59, 59
    )


def test_get_report_filters_valores_por_defecto():
    assert get_report_filters({}) == {
        "tipo": "puntos",
        "usuario": "",
        "fecha_inicio": "",
        "fecha_fin": "",
        "movimiento_tipo": "",
        "pagina": 1,
        "por_pagina": 20,
    }


def test_get_report_filters_normaliza_entradas():
    filtros = get_report_filters({
        "tipo": "  ADMIN ",
        "usuario": "  Ana  ",
        "fecha_inicio": " 2024-01-01 ",
        "movimiento_tipo": " SUMA ",
        "pagina": "3",
        "por_pagina": "5",
    })

    assert filtros["tipo"] == "admin"
    assert filtros["usuario"] == "Ana"
    assert filtros["fecha_inicio"] == "2024-01-01"
    assert filtros["movimiento_tipo"] == "suma"
    assert filtros["pagina"] == 3
    assert filtros["por_pagina"] == 5


def test_get_report_filters_pagina_no_numerica():
    with pytest.raises(ValueError):
        get_report_filters({"pagina": "abc"})


# ------------------------------------------------------------------- consultas


@pytest.fixture
def movimientos(db):
    registros = [
        MovimientoPuntos(
            paciente_id="1", idcita="10", tipo="suma", cantidad=100,
            usuario_admin="ana", fecha=datetime(2024, 1, 1, 10, 0, 0), motivo="carga",
        ),
        MovimientoPuntos(
            paciente_id="1", idcita="11", tipo="resta", cantidad=40,
            usuario_admin="beto", fecha=datetime(2024, 2, 1, 10, 0, 0), motivo="canje",
        ),
        MovimientoPuntos(
            paciente_id="1", idcita="12", tipo="suma", cantidad=500,
            usuario_admin="ana", fecha=datetime(2024, 3, 1, 10, 0, 0), motivo="reseña",
        ),
        MovimientoPuntos(
            paciente_id="2", idcita="13", tipo="suma", cantidad=7,
            usuario_admin="ana", fecha=datetime(2024, 3, 5, 10, 0, 0), motivo="carga",
        ),
    ]
    db.add_all(registros)
    db.commit()
    return registros


def test_build_filtered_movimientos_query_sin_filtros(db, movimientos):
    query = build_filtered_movimientos_query(db, get_report_filters({}))

    assert query.count() == 4


def test_build_filtered_movimientos_query_por_usuario(db, movimientos):
    query = build_filtered_movimientos_query(db, get_report_filters({"usuario": "ANA"}))

    assert query.count() == 3


def test_build_filtered_movimientos_query_por_rango_de_fechas(db, movimientos):
    filtros = get_report_filters({
        "fecha_inicio": "2024-02-01",
        "fecha_fin": "2024-03-01",
    })

    assert build_filtered_movimientos_query(db, filtros).count() == 2


def test_build_filtered_movimientos_query_por_tipo(db, movimientos):
    filtros = get_report_filters({"movimiento_tipo": "resta"})

    assert build_filtered_movimientos_query(db, filtros).count() == 1


def test_build_filtered_movimientos_query_ignora_tipo_desconocido(db, movimientos):
    filtros = get_report_filters({"movimiento_tipo": "otro"})

    assert build_filtered_movimientos_query(db, filtros).count() == 4


@pytest.fixture
def auditorias(db):
    db.add_all([
        Auditoria(
            usuario_admin="ana", accion="crear", usuario_afectado="x",
            detalle="d1", fecha=datetime(2024, 1, 1, 8, 0, 0),
        ),
        Auditoria(
            usuario_admin="beto", accion="eliminar", usuario_afectado="y",
            detalle="d2", fecha=datetime(2024, 4, 1, 8, 0, 0),
        ),
    ])
    db.commit()


def test_build_filtered_auditoria_query_filtra_usuario_y_fecha(db, auditorias):
    assert build_filtered_auditoria_query(db, get_report_filters({})).count() == 2
    assert build_filtered_auditoria_query(
        db, get_report_filters({"usuario": "beto"})
    ).count() == 1
    assert build_filtered_auditoria_query(
        db, get_report_filters({"fecha_fin": "2024-01-01"})
    ).count() == 1


def test_calcular_saldo_local_ignora_resenas(db, movimientos):
    assert calcular_saldo_local(db, "1") == 60


def test_calcular_saldo_local_sin_movimientos(db):
    assert calcular_saldo_local(db, "999") == 0


def test_ya_tiene_resena_local(db, movimientos):
    assert ya_tiene_resena_local(db, "1") is True
    assert ya_tiene_resena_local(db, "2") is False


def test_get_unread_notifications_count(db):
    db.add_all([
        Notificacion(mensaje="m1", usuario="ana", leido=0, fecha=datetime.now()),
        Notificacion(mensaje="m2", usuario="ana", leido=1, fecha=datetime.now()),
        Notificacion(mensaje="m3", usuario="beto", leido=0, fecha=datetime.now()),
    ])
    db.commit()

    assert get_unread_notifications_count(db, "ana") == 1
    assert get_unread_notifications_count(db, "") == 0


def test_crear_base_es_idempotente(db):
    crear_base(db)
    crear_base(db)

    assert db.query(Permiso).count() == 4
    assert db.query(Rol).count() == 3
    assert db.query(Rol).filter_by(nombre="superadmin").first().permisos == "all"


def test_crear_base_actualiza_rol_sin_permiso_puntos(db):
    db.add(Rol(nombre="admin", permisos="dashboard"))
    db.commit()

    crear_base(db)

    assert "puntos" in db.query(Rol).filter_by(nombre="admin").first().permisos


# -------------------------------------------------------------------- reportes


def test_build_report_context_pagina_y_kpis(db, movimientos, auditorias):
    db.add(Usuario(username="ana", rol="admin", permisos="dashboard"))
    db.add(Usuario(username="zoe", rol="usuario", permisos=""))
    db.commit()

    contexto = build_report_context(db, get_report_filters({"por_pagina": "2"}))

    assert contexto["tipo"] == "puntos"
    assert contexto["total_movimientos"] == 4
    assert len(contexto["datos_puntos"]) == 2
    assert contexto["kpis"]["usuarios"] == 2
    assert contexto["kpis"]["admins"] == 1
    assert contexto["kpis"]["puntos_netos"] == 567
    assert contexto["chart_roles"]["valores"] == [1, 1, 0]
    assert contexto["chart_principal"] == contexto["chart_movimientos_admin"]
    assert contexto["chart_fechas"] == contexto["chart_puntos_fecha"]


def test_build_report_context_tipo_admin_usa_charts_de_auditoria(db, auditorias):
    contexto = build_report_context(db, get_report_filters({"tipo": "admin"}))

    assert contexto["chart_principal"] == contexto["chart_auditoria"]
    assert contexto["chart_fechas"] == contexto["chart_auditoria_fecha"]
    assert contexto["kpis"]["puntos_netos"] == 2
    assert contexto["kpis"]["movimientos_hoy"] == 2
    assert sorted(contexto["chart_auditoria"]["labels"]) == ["crear", "eliminar"]


def test_build_report_context_segunda_pagina(db, movimientos):
    contexto = build_report_context(db, get_report_filters({"pagina": "2", "por_pagina": "3"}))

    assert contexto["pagina"] == 2
    assert len(contexto["datos_puntos"]) == 1


@pytest.fixture
def contexto_reporte(db, movimientos, auditorias):
    return build_report_context(db, get_report_filters({}))


def test_exportar_reportes_csv(flask_app, contexto_reporte):
    with flask_app.test_request_context("/"):
        response = exportar_reportes_csv(contexto_reporte)

    cuerpo = response.get_data(as_text=True)
    lineas = [linea for linea in cuerpo.splitlines() if linea]

    assert response.mimetype == "text/csv"
    assert "reporte_sistema.csv" in response.headers["Content-Disposition"]
    assert lineas[0].startswith("seccion,")
    assert len([l for l in lineas if l.startswith("puntos,")]) == 4
    assert len([l for l in lineas if l.startswith("auditoria,")]) == 2


def test_exportar_reportes_excel(flask_app, contexto_reporte):
    import io

    from openpyxl import load_workbook

    with flask_app.test_request_context("/"):
        response = exportar_reportes_excel(contexto_reporte)

    workbook = load_workbook(io.BytesIO(response.get_data()))

    assert workbook.sheetnames == ["Resumen", "Puntos", "Auditoria"]
    assert workbook["Resumen"]["A2"].value == "Usuarios"
    assert workbook["Puntos"].max_row == 5
    assert workbook["Auditoria"].max_row == 3
    assert "reporte_sistema.xlsx" in response.headers["Content-Disposition"]


# ------------------------------------------------------------- sesión / acceso


def test_get_usuario_lee_la_sesion(flask_app):
    with flask_app.test_request_context("/"):
        assert get_usuario() is None
        session["usuario"] = "ana"
        assert get_usuario() == "ana"


@pytest.mark.parametrize("rol, permisos, esperado", [
    ("superadmin", [], "/dashboard"),
    ("admin", ["dashboard", "puntos"], "/dashboard"),
    ("admin", ["puntos"], "/puntos"),
    ("admin", ["usuarios"], "/usuarios"),
    ("admin", ["reportes"], "/reportes"),
    ("admin", ["transcripcion"], "/transcripcion"),
    ("usuario", [], "/logout"),
])
def test_obtener_pagina_inicio(flask_app, rol, permisos, esperado):
    with flask_app.test_request_context("/"):
        session["rol"] = rol
        session["permisos"] = permisos

        assert obtener_pagina_inicio() == esperado


def test_get_db_reutiliza_sesion(flask_app, monkeypatch, db):
    monkeypatch.setattr(app_module, "SessionLocal", lambda: db)

    with flask_app.test_request_context("/"):
        primera = get_db()

        assert primera is db
        assert get_db() is primera


def test_log_auditoria_persiste_registro(request_ctx, db):
    log_auditoria("ana", "crear", "beto", "detalle")

    registro = db.query(Auditoria).one()

    assert (registro.usuario_admin, registro.accion) == ("ana", "crear")
    assert registro.usuario_afectado == "beto"
    assert registro.fecha is not None


def _vista_protegida():
    @requiere_permiso("puntos")
    def vista():
        return "ok"

    return vista


def test_requiere_permiso_redirige_sin_sesion(request_ctx):
    respuesta = _vista_protegida()()

    assert respuesta.status_code == 302


def test_requiere_permiso_redirige_usuario_inexistente(request_ctx):
    session["usuario"] = "fantasma"

    assert _vista_protegida()().status_code == 302


def test_requiere_permiso_superadmin_siempre_pasa(request_ctx, db):
    db.add(Usuario(username="root", rol="superadmin", permisos=""))
    db.commit()
    session["usuario"] = "root"

    assert _vista_protegida()() == "ok"


def test_requiere_permiso_usa_permisos_del_rol(request_ctx, db):
    db.add(Usuario(username="ana", rol="admin", permisos=""))
    db.add(Rol(nombre="admin", permisos="dashboard,puntos"))
    db.commit()
    session["usuario"] = "ana"

    assert _vista_protegida()() == "ok"


def test_requiere_permiso_usa_permisos_del_usuario(request_ctx, db):
    db.add(Usuario(username="ana", rol="usuario", permisos="puntos"))
    db.commit()
    session["usuario"] = "ana"

    assert _vista_protegida()() == "ok"


def test_requiere_permiso_deniega_sin_permiso(request_ctx, db):
    db.add(Usuario(username="ana", rol="usuario", permisos="dashboard"))
    db.commit()
    session["usuario"] = "ana"

    assert _vista_protegida()() == ("No tienes permiso", 403)


def test_requiere_permiso_conserva_nombre():
    @requiere_permiso("puntos")
    def vista_puntos():
        return "ok"

    assert vista_puntos.__name__ == "vista_puntos"
